from docx import Document


def read_docx_blocks(file_path: str):
    doc = Document(file_path)
    blocks = []

    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if text:
            blocks.append({
                "type": "paragraph",
                "text": text,
                "style": getattr(p.style, "name", "") if p.style else "",
            })

    for table in doc.tables:
        rows_text = []
        for row in table.rows:
            cols = []
            for cell in row.cells:
                cell_text = " ".join([para.text.strip() for para in cell.paragraphs if para.text.strip()])
                if cell_text:
                    cols.append(cell_text)
            if cols:
                rows_text.append(" | ".join(cols))
        if rows_text:
            blocks.append({
                "type": "table",
                "text": "\n".join(rows_text),
                "style": "table",
            })

    return blocks


def blocks_to_raw_text(blocks):
    return "\n".join([b["text"] for b in blocks if b.get("text")])
