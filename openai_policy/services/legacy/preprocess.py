from __future__ import annotations

import io
import json
import os
import re
import shutil
import unicodedata
import uuid
from collections import OrderedDict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .compat import APIRouter, BackgroundTasks, File, Form, Header, HTTPException, UploadFile
from .compat import FileResponse, PlainTextResponse
from .compat import OpenAI
from .compat import BaseModel, Field

from .admin import (
    _delete_vector_file,
    _get_store_id,
    _iter_vector_files,
    _require_admin,
    _safe_token,
    _serialize_vector_file,
    openai_client,
)

try:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
except Exception:  # pragma: no cover - dependency is declared in requirements.txt
    Document = None
    Table = None
    Paragraph = None
    CT_Tbl = None
    CT_P = None


router = APIRouter(prefix="/api/admin", tags=["Policy Preprocess"])

BASE_DIR = Path(__file__).resolve().parent.parent
WORK_DIR = Path(os.environ.get("POLICY_WORK_DIR", str(BASE_DIR / "data"))).resolve()
PREPROCESS_DIR = WORK_DIR / "preprocess"
EFFECTIVE_DIR = WORK_DIR / "effective"
PREPROCESS_DIR.mkdir(parents=True, exist_ok=True)
EFFECTIVE_DIR.mkdir(parents=True, exist_ok=True)

OPENAI_PREPROCESS_MODEL = os.environ.get(
    "OPENAI_PREPROCESS_MODEL", os.environ.get("OPENAI_MODEL", "")
).strip()
PREPROCESS_MAX_UPLOAD_MB = int(os.environ.get("PREPROCESS_MAX_UPLOAD_MB", "50"))
PREPROCESS_BATCH_CHARS = int(os.environ.get("PREPROCESS_BATCH_CHARS", "12000"))
PREPROCESS_ALLOWED_EXTENSIONS = {".docx", ".md", ".txt"}

LEGACY_FONT_MARKERS = (
    "sutonny",
    "bijoy",
    "boishakhi",
    "proshika",
    "bangla2000",
    "bangsee",
)
LEGACY_GLYPH_MARKERS = set("‡†©ª¯šœ™¤¦µ¶¬")
BANGLA_RE = re.compile(r"[\u0980-\u09FF]")
NUMBER_RE = re.compile(r"[0-9০-৯]+(?:[./:\-][0-9০-৯]+)*(?:%|%)?")
CLAUSE_AT_START_RE = re.compile(
    r"^\s*(?:ধারা\s*)?([0-9০-৯]+(?:\s*[.।]\s*[0-9০-৯]+){0,6})\s*(?:[:\-–—.)।]|\s)",
    re.IGNORECASE,
)
CLAUSE_ANYWHERE_RE = re.compile(
    r"(?:ধারা|clause|section)\s*[:#-]?\s*([0-9০-৯]+(?:\s*[.।]\s*[0-9০-৯]+){0,6})",
    re.IGNORECASE,
)


class ApproveRequest(BaseModel):
    approved: bool = True
    note: str = Field(default="", max_length=1000)


class BuildEffectiveRequest(BaseModel):
    job_ids: list[str] = Field(min_length=1)
    output_version: str = Field(default="effective-current", max_length=80)
    title: str = Field(default="Organization Policy — Effective", max_length=200)


class PublishEffectiveRequest(BaseModel):
    replace_active_effective: bool = True


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def _atomic_write_json(path: Path, payload: dict[str, Any]) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def _job_dir(job_id: str) -> Path:
    return PREPROCESS_DIR / job_id


def _build_dir(build_id: str) -> Path:
    return EFFECTIVE_DIR / build_id


def _load_job(job_id: str) -> dict[str, Any]:
    path = _job_dir(job_id) / "job.json"
    if not path.exists():
        raise HTTPException(status_code=404, detail="Preprocess job not found.")
    return json.loads(path.read_text(encoding="utf-8"))


def _save_job(job: dict[str, Any]) -> None:
    job["updated_at"] = _utc_now()
    _atomic_write_json(_job_dir(job["job_id"]) / "job.json", job)


def _load_build(build_id: str) -> dict[str, Any]:
    path = _build_dir(build_id) / "build.json"
    if not path.exists():
        raise HTTPException(status_code=404, detail="Effective build not found.")
    return json.loads(path.read_text(encoding="utf-8"))


def _save_build(build: dict[str, Any]) -> None:
    build["updated_at"] = _utc_now()
    _atomic_write_json(_build_dir(build["build_id"]) / "build.json", build)


def _clean_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text or "")
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _docx_blocks(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    if Document is None:
        raise RuntimeError("python-docx is not installed.")

    doc = Document(io.BytesIO(data))
    blocks: list[dict[str, Any]] = []
    fonts: set[str] = set()
    index = 0

    def add_paragraph(p: Any, kind: str = "paragraph") -> None:
        nonlocal index
        text = _clean_text(getattr(p, "text", "") or "")
        for run in getattr(p, "runs", []) or []:
            name = getattr(getattr(run, "font", None), "name", None)
            if name:
                fonts.add(str(name))
        if not text:
            return
        index += 1
        style = ""
        try:
            style = str(getattr(getattr(p, "style", None), "name", "") or "")
        except Exception:
            style = ""
        blocks.append({"id": index, "kind": kind, "style": style, "text": text})

    for child in doc.element.body.iterchildren():
        if CT_P is not None and isinstance(child, CT_P):
            add_paragraph(Paragraph(child, doc))
        elif CT_Tbl is not None and isinstance(child, CT_Tbl):
            table = Table(child, doc)
            for r_idx, row in enumerate(table.rows, start=1):
                cells: list[str] = []
                for cell in row.cells:
                    cell_texts: list[str] = []
                    for p in cell.paragraphs:
                        for run in p.runs:
                            name = getattr(run.font, "name", None)
                            if name:
                                fonts.add(str(name))
                        txt = _clean_text(p.text)
                        if txt:
                            cell_texts.append(txt)
                    cells.append(" ".join(cell_texts))
                row_text = " | ".join(cells).strip(" |")
                if row_text:
                    index += 1
                    blocks.append(
                        {
                            "id": index,
                            "kind": "table_row",
                            "style": f"Table row {r_idx}",
                            "text": row_text,
                        }
                    )

    return blocks, sorted(fonts)


def _text_blocks(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    text = ""
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    blocks: list[dict[str, Any]] = []
    for idx, line in enumerate(text.splitlines(), start=1):
        line = _clean_text(line)
        if line:
            blocks.append({"id": idx, "kind": "paragraph", "style": "", "text": line})
    return blocks, []


def _extract_blocks(data: bytes, suffix: str) -> tuple[list[dict[str, Any]], list[str]]:
    if suffix == ".docx":
        return _docx_blocks(data)
    return _text_blocks(data)


def _legacy_report(blocks: list[dict[str, Any]], fonts: list[str]) -> dict[str, Any]:
    text = "\n".join(b["text"] for b in blocks)
    bangla_count = len(BANGLA_RE.findall(text))
    latin_count = len(re.findall(r"[A-Za-z]", text))
    suspicious_glyph_count = sum(1 for ch in text if ch in LEGACY_GLYPH_MARKERS)
    legacy_fonts = [f for f in fonts if any(marker in f.lower() for marker in LEGACY_FONT_MARKERS)]
    risk = bool(legacy_fonts or suspicious_glyph_count >= 3)
    if bangla_count == 0 and latin_count > 100 and suspicious_glyph_count >= 1:
        risk = True
    return {
        "bangla_chars": bangla_count,
        "latin_chars": latin_count,
        "suspicious_legacy_glyphs": suspicious_glyph_count,
        "fonts": fonts,
        "legacy_fonts": legacy_fonts,
        "legacy_risk": risk,
    }


def _make_batches(blocks: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    batches: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    chars = 0
    for block in blocks:
        size = len(block["text"]) + 80
        if current and chars + size > PREPROCESS_BATCH_CHARS:
            batches.append(current)
            current = []
            chars = 0
        current.append(block)
        chars += size
    if current:
        batches.append(current)
    return batches


def _normalize_batch_with_openai(batch: list[dict[str, Any]], model: str) -> tuple[list[dict[str, Any]], dict[str, int]]:
    schema = {
        "type": "object",
        "properties": {
            "items": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "id": {"type": "integer"},
                        "text": {"type": "string"},
                        "changed": {"type": "boolean"},
                        "review_required": {"type": "boolean"},
                        "reason": {"type": "string"},
                    },
                    "required": ["id", "text", "changed", "review_required", "reason"],
                    "additionalProperties": False,
                },
            }
        },
        "required": ["items"],
        "additionalProperties": False,
    }
    payload = [{"id": b["id"], "text": b["text"]} for b in batch]
    instructions = (
        "You are a LOSSLESS Bengali policy-document normalization engine. "
        "For every input item return exactly one item with the same id. Never summarize, translate, paraphrase, "
        "interpret, or improve policy language. Never change policy meaning, numbers, dates, percentages, names, "
        "clause references, approvals, limits, or conditions. If text is already proper Unicode Bangla/English, "
        "preserve the wording and only normalize Unicode/spacing/punctuation when necessary. If text is legacy "
        "Bijoy/SutonnyMJ-style Bangla represented as Latin/ANSI glyph sequences, convert it to the equivalent "
        "Unicode Bangla while preserving wording exactly. If any conversion is uncertain, keep the safest text "
        "possible and set review_required=true with a concise reason. Do not invent missing words."
    )
    response = openai_client.responses.create(
        model=model,
        instructions=instructions,
        input=json.dumps(payload, ensure_ascii=False),
        text={
            "format": {
                "type": "json_schema",
                "name": "policy_unicode_normalization",
                "strict": True,
                "schema": schema,
            }
        },
        store=False,
    )
    raw = (getattr(response, "output_text", "") or "").strip()
    if not raw:
        raise RuntimeError("OpenAI preprocessing returned an empty response.")
    parsed = json.loads(raw)
    usage = getattr(response, "usage", None)
    usage_dict = {
        "input_tokens": int(getattr(usage, "input_tokens", 0) or 0),
        "output_tokens": int(getattr(usage, "output_tokens", 0) or 0),
        "total_tokens": int(getattr(usage, "total_tokens", 0) or 0),
    }
    return list(parsed.get("items", [])), usage_dict


def _normalize_blocks(
    blocks: list[dict[str, Any]],
    *,
    use_openai: bool,
    model: str,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    normalized = [dict(b, text=_clean_text(b["text"])) for b in blocks]
    report: dict[str, Any] = {
        "used_openai": False,
        "model": model if use_openai else "",
        "changed_blocks": 0,
        "review_required_blocks": 0,
        "review_items": [],
        "usage": {"input_tokens": 0, "output_tokens": 0, "total_tokens": 0},
    }
    if not use_openai:
        return normalized, report
    if not model:
        raise RuntimeError("OPENAI_PREPROCESS_MODEL or OPENAI_MODEL is not configured.")

    by_id = {int(b["id"]): b for b in normalized}
    report["used_openai"] = True

    for batch in _make_batches(normalized):
        output_items, usage = _normalize_batch_with_openai(batch, model)
        for k in report["usage"]:
            report["usage"][k] += usage.get(k, 0)

        returned: set[int] = set()
        for item in output_items:
            item_id = int(item.get("id", -1))
            if item_id not in by_id:
                continue
            returned.add(item_id)
            original = by_id[item_id]["text"]
            new_text = _clean_text(str(item.get("text", "") or ""))
            review_required = bool(item.get("review_required", False))
            reason = str(item.get("reason", "") or "").strip()

            original_numbers = NUMBER_RE.findall(original)
            new_numbers = NUMBER_RE.findall(new_text)
            if original_numbers != new_numbers:
                review_required = True
                reason = (reason + "; " if reason else "") + "numeric/date tokens changed"

            if not new_text:
                new_text = original
                review_required = True
                reason = (reason + "; " if reason else "") + "empty normalized text"

            if new_text != original:
                report["changed_blocks"] += 1
            if review_required:
                report["review_required_blocks"] += 1
                report["review_items"].append(
                    {"id": item_id, "reason": reason or "manual review requested", "before": original[:500], "after": new_text[:500]}
                )

            by_id[item_id]["text"] = new_text

        missing = [int(b["id"]) for b in batch if int(b["id"]) not in returned]
        for item_id in missing:
            report["review_required_blocks"] += 1
            report["review_items"].append(
                {"id": item_id, "reason": "OpenAI did not return this source block; original kept", "before": by_id[item_id]["text"][:500], "after": by_id[item_id]["text"][:500]}
            )

    return [by_id[int(b["id"])] for b in normalized], report


def _normalize_clause_id(value: str) -> str:
    value = re.sub(r"\s+", "", value or "")
    value = value.replace("।", ".")
    return value.strip(".:-–—() ")


def _detect_clause_id(text: str) -> str:
    match = CLAUSE_AT_START_RE.search(text)
    if not match:
        match = CLAUSE_ANYWHERE_RE.search(text[:300])
    return _normalize_clause_id(match.group(1)) if match else ""


def _is_heading(block: dict[str, Any], clause_id: str) -> bool:
    style = str(block.get("style", "") or "").lower()
    text = block["text"]
    return bool(clause_id or "heading" in style or (len(text) <= 120 and text.endswith(":")))


def _atomic_records(blocks: list[dict[str, Any]], job: dict[str, Any]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    synthetic = 0

    def flush() -> None:
        nonlocal current
        if not current:
            return
        current["text"] = "\n".join(current.pop("parts")).strip()
        records.append(current)
        current = None

    for block in blocks:
        text = block["text"].strip()
        if not text:
            continue
        clause_id = _detect_clause_id(text)
        if _is_heading(block, clause_id):
            flush()
            synthetic += 1
            current = {
                "record_id": f"C-{clause_id}" if clause_id else f"H-{synthetic:04d}",
                "clause_id": clause_id,
                "heading": text if len(text) <= 250 else "",
                "record_type": "clause" if clause_id else "section",
                "source_start": block["id"],
                "source_end": block["id"],
                "parts": [text],
            }
        elif current is not None:
            current["parts"].append(text)
            current["source_end"] = block["id"]
        else:
            synthetic += 1
            records.append(
                {
                    "record_id": f"P-{synthetic:04d}",
                    "clause_id": clause_id,
                    "heading": "",
                    "record_type": "paragraph" if block["kind"] != "table_row" else "table_row",
                    "source_start": block["id"],
                    "source_end": block["id"],
                    "text": text,
                }
            )
    flush()

    for rec in records:
        rec.update(
            {
                "document_key": job["document_key"],
                "document_type": job["document_type"],
                "version": job["version"],
                "effective_date": job["effective_date"],
                "source_filename": job["original_filename"],
                "amendment_action": (
                    "replace_or_add"
                    if job["document_type"] == "amendment" and rec.get("clause_id")
                    else "append"
                ),
            }
        )
    return records


def _write_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as fh:
        for rec in records:
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")


def _doc_markdown(job: dict[str, Any], records: list[dict[str, Any]]) -> str:
    lines = [
        "---",
        f"document_key: {job['document_key']}",
        f"document_type: {job['document_type']}",
        f"version: {job['version']}",
        f"effective_date: {job['effective_date'] or 'unspecified'}",
        f"source_file: {job['original_filename']}",
        "pipeline: docx-unicode-atomic-json-markdown",
        "---",
        "",
        f"# {job['title'] or job['original_filename']}",
        "",
    ]
    for rec in records:
        clause_id = rec.get("clause_id", "")
        if clause_id:
            lines.extend(
                [
                    f"## Clause {clause_id}",
                    f"Source-Type: {job['document_type']} | Version: {job['version']} | Effective-Date: {job['effective_date'] or 'unspecified'}",
                    "",
                    rec["text"].strip(),
                    "",
                ]
            )
        else:
            lines.extend([rec["text"].strip(), ""])
    return "\n".join(lines).strip() + "\n"


def _run_preprocess(job_id: str, use_openai: bool) -> None:
    job = _load_job(job_id)
    work = _job_dir(job_id)
    try:
        job["status"] = "extracting"
        _save_job(job)
        source = work / job["source_artifact"]
        data = source.read_bytes()
        suffix = Path(job["original_filename"]).suffix.lower()
        blocks, fonts = _extract_blocks(data, suffix)
        if not blocks:
            raise RuntimeError("No readable text was extracted from the source document.")

        extracted_text = "\n".join(b["text"] for b in blocks).strip() + "\n"
        (work / "extracted.txt").write_text(extracted_text, encoding="utf-8")
        legacy = _legacy_report(blocks, fonts)

        job["status"] = "normalizing"
        job["validation"] = {"extraction": {"blocks": len(blocks), **legacy}}
        _save_job(job)

        normalized_blocks, ai_report = _normalize_blocks(
            blocks,
            use_openai=use_openai,
            model=OPENAI_PREPROCESS_MODEL,
        )
        unicode_text = "\n".join(b["text"] for b in normalized_blocks).strip() + "\n"
        (work / "unicode.txt").write_text(unicode_text, encoding="utf-8")

        unicode_report = _legacy_report(normalized_blocks, fonts)
        records = _atomic_records(normalized_blocks, job)
        _write_jsonl(work / "atomic.jsonl", records)
        (work / "rag.md").write_text(_doc_markdown(job, records), encoding="utf-8")

        review_required = ai_report["review_required_blocks"] > 0
        # Legacy font is a source warning, but successful Unicode conversion can still be reviewable.
        if legacy["legacy_risk"] and unicode_report["bangla_chars"] == 0:
            review_required = True
            ai_report["review_items"].append(
                {
                    "id": 0,
                    "reason": "Legacy Bangla risk detected but normalized output contains no Bangla Unicode characters.",
                    "before": "",
                    "after": "",
                }
            )

        job["status"] = "review_required" if review_required else "ready"
        job["approved"] = False
        job["validation"] = {
            "extraction": {"blocks": len(blocks), **legacy},
            "unicode_output": unicode_report,
            "openai_normalization": ai_report,
            "atomic_records": len(records),
        }
        job["artifacts"] = {
            "extracted": "extracted.txt",
            "unicode": "unicode.txt",
            "atomic": "atomic.jsonl",
            "markdown": "rag.md",
        }
        _save_job(job)
    except Exception as exc:
        job["status"] = "failed"
        job["error"] = str(exc)
        _save_job(job)


def _list_jobs() -> list[dict[str, Any]]:
    jobs: list[dict[str, Any]] = []
    for path in PREPROCESS_DIR.glob("*/job.json"):
        try:
            jobs.append(json.loads(path.read_text(encoding="utf-8")))
        except Exception:
            continue
    jobs.sort(key=lambda x: x.get("created_at", ""), reverse=True)
    return jobs


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            records.append(json.loads(line))
    return records


def _version_sort_key(version: str) -> tuple[Any, ...]:
    parts = re.split(r"(\d+)", version or "")
    return tuple(int(p) if p.isdigit() else p.lower() for p in parts)


def _job_sort_key(job: dict[str, Any]) -> tuple[Any, ...]:
    type_rank = 0 if job.get("document_type") == "base" else 1
    return (
        type_rank,
        job.get("effective_date") or "0000-00-00",
        _version_sort_key(job.get("version", "")),
        job.get("created_at", ""),
    )


def _effective_markdown(build: dict[str, Any], ordered_jobs: list[dict[str, Any]], records: list[dict[str, Any]]) -> str:
    lines = [
        "---",
        f"document_key: {build['document_key']}",
        "document_type: effective",
        f"version: {build['output_version']}",
        f"generated_at: {build['created_at']}",
        "pipeline: base-amendments-effective-markdown",
        "---",
        "",
        f"# {build['title']}",
        "",
        "## Source Version Chain",
    ]
    for job in ordered_jobs:
        lines.append(
            f"- {job['document_type']} | {job['version']} | effective {job['effective_date'] or 'unspecified'} | {job['original_filename']}"
        )
    lines.extend(["", "## Effective Clauses", ""])

    for rec in records:
        clause_id = rec.get("clause_id", "")
        src = rec.get("effective_source", {})
        if clause_id:
            lines.extend(
                [
                    f"### Clause {clause_id}",
                    (
                        f"Effective-Source: {src.get('document_type', '')} | Version: {src.get('version', '')} "
                        f"| Effective-Date: {src.get('effective_date') or 'unspecified'} | File: {src.get('source_filename', '')}"
                    ),
                    "",
                    rec["text"].strip(),
                    "",
                ]
            )
        else:
            lines.extend(
                [
                    "### Additional Effective Provision",
                    (
                        f"Effective-Source: {src.get('document_type', '')} | Version: {src.get('version', '')} "
                        f"| Effective-Date: {src.get('effective_date') or 'unspecified'} | File: {src.get('source_filename', '')}"
                    ),
                    "",
                    rec["text"].strip(),
                    "",
                ]
            )
    return "\n".join(lines).strip() + "\n"


def _upsert_effective_bytes(build: dict[str, Any], data: bytes, replace_active: bool) -> dict[str, Any]:
    vector_store_id = _get_store_id()
    document_key = build["document_key"]
    version = build["output_version"]
    latest_effective_date = build.get("latest_effective_date", "") or "unspecified"
    filename = f"{_safe_token(document_key)}__effective__{_safe_token(version, 'current')}__{latest_effective_date}.md"
    attrs: dict[str, str | int | bool | float] = {
        "document_key": document_key,
        "document_type": "effective",
        "version": version,
        "effective_date": latest_effective_date,
        "title": build["title"],
        "active": True,
        "source_build_id": build["build_id"],
        "pipeline": "docx-unicode-json-effective-md",
    }

    old_matches: list[Any] = []
    if replace_active:
        for item in _iter_vector_files(vector_store_id):
            a = dict(getattr(item, "attributes", None) or {})
            if str(a.get("document_key", "")) == document_key and str(a.get("document_type", "")) == "effective":
                old_matches.append(item)

    uploaded_file_id = ""
    attached_file_id = ""
    removed: list[str] = []
    warnings: list[str] = []
    try:
        uploaded = openai_client.files.create(
            file=(filename, data, "text/markdown"),
            purpose="assistants",
        )
        uploaded_file_id = str(uploaded.id)
        attached = openai_client.vector_stores.files.create(
            vector_store_id,
            file_id=uploaded.id,
            attributes=attrs,
        )
        attached_file_id = str(attached.id)
        indexed = openai_client.vector_stores.files.poll(
            attached.id,
            vector_store_id=vector_store_id,
        )
        status = str(getattr(indexed, "status", "") or "")
        if status != "completed":
            last_error = getattr(indexed, "last_error", None)
            msg = str(getattr(last_error, "message", "") or last_error or status)
            try:
                _delete_vector_file(vector_store_id, attached_file_id, delete_file=True)
            except Exception:
                pass
            raise RuntimeError(f"Vector indexing failed: {msg}")

        for old in old_matches:
            old_id = str(getattr(old, "id", "") or "")
            if not old_id or old_id == attached_file_id:
                continue
            try:
                _delete_vector_file(vector_store_id, old_id, delete_file=True)
                removed.append(old_id)
            except Exception as exc:
                warnings.append(f"Could not remove previous effective file {old_id}: {exc}")

        return {
            "ok": True,
            "action": "replaced_active_effective" if removed else "inserted_effective",
            "removed_file_ids": removed,
            "warnings": warnings,
            "file": _serialize_vector_file(indexed),
        }
    except Exception:
        if attached_file_id:
            try:
                _delete_vector_file(vector_store_id, attached_file_id, delete_file=True)
            except Exception:
                pass
        elif uploaded_file_id:
            try:
                openai_client.files.delete(uploaded_file_id)
            except Exception:
                pass
        raise


@router.get("/preprocess/config")
def preprocess_config(x_admin_token: str | None = Header(default=None)) -> dict[str, Any]:
    _require_admin(x_admin_token)
    return {
        "model": OPENAI_PREPROCESS_MODEL,
        "max_upload_mb": PREPROCESS_MAX_UPLOAD_MB,
        "allowed_extensions": sorted(PREPROCESS_ALLOWED_EXTENSIONS),
        "work_dir": str(WORK_DIR),
        "pipeline": ["extract", "unicode", "atomic_jsonl", "rag_markdown", "approve", "build_effective", "upsert"],
    }


@router.post("/preprocess", status_code=202)
def create_preprocess_job(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    document_key: str = Form(...),
    document_type: str = Form(...),
    version: str = Form(...),
    effective_date: str = Form(""),
    title: str = Form(""),
    use_openai_normalization: bool = Form(True),
    x_admin_token: str | None = Header(default=None),
) -> dict[str, Any]:
    _require_admin(x_admin_token)
    original_filename = Path(file.filename or "policy.docx").name
    suffix = Path(original_filename).suffix.lower()
    if suffix not in PREPROCESS_ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Preprocess supports: {', '.join(sorted(PREPROCESS_ALLOWED_EXTENSIONS))}")

    document_type = (document_type or "").strip().lower()
    if document_type not in {"base", "amendment"}:
        raise HTTPException(status_code=400, detail="document_type must be 'base' or 'amendment'.")
    version = (version or "").strip()
    if not version:
        raise HTTPException(status_code=400, detail="version is required.")
    effective_date = (effective_date or "").strip()
    if effective_date:
        try:
            datetime.strptime(effective_date, "%Y-%m-%d")
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="effective_date must be YYYY-MM-DD.") from exc

    data = file.file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    if len(data) > PREPROCESS_MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"File exceeds {PREPROCESS_MAX_UPLOAD_MB} MB.")

    job_id = f"pp_{uuid.uuid4().hex[:16]}"
    work = _job_dir(job_id)
    work.mkdir(parents=True, exist_ok=False)
    source_name = "source" + suffix
    (work / source_name).write_bytes(data)
    job = {
        "job_id": job_id,
        "status": "queued",
        "approved": False,
        "approval_note": "",
        "document_key": _safe_token(document_key, "organization-policy"),
        "document_type": document_type,
        "version": version[:80],
        "effective_date": effective_date,
        "title": (title or original_filename)[:200],
        "original_filename": original_filename,
        "source_artifact": source_name,
        "source_bytes": len(data),
        "use_openai_normalization": bool(use_openai_normalization),
        "created_at": _utc_now(),
        "updated_at": _utc_now(),
        "validation": {},
        "artifacts": {},
        "error": "",
    }
    _save_job(job)
    background_tasks.add_task(_run_preprocess, job_id, bool(use_openai_normalization))
    return {"ok": True, "job_id": job_id, "status": "queued", "poll_url": f"/api/admin/preprocess/{job_id}"}


@router.get("/preprocess/jobs")
def list_preprocess_jobs(x_admin_token: str | None = Header(default=None)) -> dict[str, Any]:
    _require_admin(x_admin_token)
    jobs = _list_jobs()
    return {"count": len(jobs), "jobs": jobs}


@router.get("/preprocess/{job_id}")
def get_preprocess_job(job_id: str, x_admin_token: str | None = Header(default=None)) -> dict[str, Any]:
    _require_admin(x_admin_token)
    return _load_job(job_id)


@router.post("/preprocess/{job_id}/approve")
def approve_preprocess_job(
    job_id: str,
    request: ApproveRequest,
    x_admin_token: str | None = Header(default=None),
) -> dict[str, Any]:
    _require_admin(x_admin_token)
    job = _load_job(job_id)
    if job["status"] not in {"ready", "review_required", "approved"}:
        raise HTTPException(status_code=409, detail=f"Job cannot be approved while status={job['status']}.")
    job["approved"] = bool(request.approved)
    job["approval_note"] = request.note
    job["approved_at"] = _utc_now() if request.approved else ""
    if request.approved:
        job["status"] = "approved"
    elif job["status"] == "approved":
        job["status"] = "ready"
    _save_job(job)
    return {"ok": True, "job": job}


@router.get("/preprocess/{job_id}/artifact/{artifact}")
def get_preprocess_artifact(
    job_id: str,
    artifact: str,
    download: bool = False,
    x_admin_token: str | None = Header(default=None),
):
    _require_admin(x_admin_token)
    job = _load_job(job_id)
    mapping = {
        "source": job.get("source_artifact", ""),
        "extracted": job.get("artifacts", {}).get("extracted", ""),
        "unicode": job.get("artifacts", {}).get("unicode", ""),
        "atomic": job.get("artifacts", {}).get("atomic", ""),
        "markdown": job.get("artifacts", {}).get("markdown", ""),
    }
    filename = mapping.get(artifact, "")
    if not filename:
        raise HTTPException(status_code=404, detail="Artifact is not available yet.")
    path = _job_dir(job_id) / filename
    if not path.exists():
        raise HTTPException(status_code=404, detail="Artifact file not found.")
    if download or artifact == "source":
        return FileResponse(path, filename=path.name)
    return PlainTextResponse(path.read_text(encoding="utf-8"), media_type="text/plain; charset=utf-8")


@router.delete("/preprocess/{job_id}")
def delete_preprocess_job(job_id: str, x_admin_token: str | None = Header(default=None)) -> dict[str, Any]:
    _require_admin(x_admin_token)
    work = _job_dir(job_id)
    if not work.exists():
        raise HTTPException(status_code=404, detail="Preprocess job not found.")
    shutil.rmtree(work)
    return {"ok": True, "job_id": job_id}


@router.post("/effective/build")
def build_effective_policy(
    request: BuildEffectiveRequest,
    x_admin_token: str | None = Header(default=None),
) -> dict[str, Any]:
    _require_admin(x_admin_token)
    jobs: list[dict[str, Any]] = []
    seen: set[str] = set()
    for job_id in request.job_ids:
        if job_id in seen:
            continue
        seen.add(job_id)
        job = _load_job(job_id)
        if not job.get("approved"):
            raise HTTPException(status_code=409, detail=f"Job {job_id} is not approved.")
        if not (_job_dir(job_id) / "atomic.jsonl").exists():
            raise HTTPException(status_code=409, detail=f"Job {job_id} has no atomic JSON artifact.")
        jobs.append(job)

    keys = {j["document_key"] for j in jobs}
    if len(keys) != 1:
        raise HTTPException(status_code=400, detail="All selected jobs must use the same document_key.")
    if not any(j["document_type"] == "base" for j in jobs):
        raise HTTPException(status_code=400, detail="Select at least one approved base policy job.")

    ordered_jobs = sorted(jobs, key=_job_sort_key)
    clause_map: "OrderedDict[str, dict[str, Any]]" = OrderedDict()
    appendix: list[dict[str, Any]] = []
    replacement_log: list[dict[str, str]] = []

    for job in ordered_jobs:
        records = _read_jsonl(_job_dir(job["job_id"]) / "atomic.jsonl")
        for rec in records:
            rec = dict(rec)
            rec["effective_source"] = {
                "job_id": job["job_id"],
                "document_type": job["document_type"],
                "version": job["version"],
                "effective_date": job["effective_date"],
                "source_filename": job["original_filename"],
            }
            cid = str(rec.get("clause_id", "") or "")
            if cid:
                if cid in clause_map:
                    old = clause_map[cid].get("effective_source", {})
                    replacement_log.append(
                        {
                            "clause_id": cid,
                            "from": f"{old.get('document_type','')} {old.get('version','')}",
                            "to": f"{job['document_type']} {job['version']}",
                        }
                    )
                clause_map[cid] = rec
            else:
                appendix.append(rec)

    effective_records = list(clause_map.values()) + appendix
    build_id = f"eff_{uuid.uuid4().hex[:16]}"
    work = _build_dir(build_id)
    work.mkdir(parents=True, exist_ok=False)
    latest_date = max((j.get("effective_date") or "" for j in ordered_jobs), default="")
    build = {
        "build_id": build_id,
        "status": "ready",
        "document_key": ordered_jobs[0]["document_key"],
        "output_version": request.output_version.strip() or "effective-current",
        "title": request.title.strip() or "Organization Policy — Effective",
        "job_ids": [j["job_id"] for j in ordered_jobs],
        "merge_order": [
            {
                "job_id": j["job_id"],
                "document_type": j["document_type"],
                "version": j["version"],
                "effective_date": j["effective_date"],
                "filename": j["original_filename"],
            }
            for j in ordered_jobs
        ],
        "latest_effective_date": latest_date,
        "effective_records": len(effective_records),
        "replacement_count": len(replacement_log),
        "replacement_log": replacement_log,
        "created_at": _utc_now(),
        "updated_at": _utc_now(),
        "published": False,
        "published_file": None,
    }
    _write_jsonl(work / "effective.jsonl", effective_records)
    md = _effective_markdown(build, ordered_jobs, effective_records)
    (work / "effective.md").write_text(md, encoding="utf-8")
    _save_build(build)
    return {"ok": True, "build": build}


@router.get("/effective/builds")
def list_effective_builds(x_admin_token: str | None = Header(default=None)) -> dict[str, Any]:
    _require_admin(x_admin_token)
    builds: list[dict[str, Any]] = []
    for path in EFFECTIVE_DIR.glob("*/build.json"):
        try:
            builds.append(json.loads(path.read_text(encoding="utf-8")))
        except Exception:
            continue
    builds.sort(key=lambda x: x.get("created_at", ""), reverse=True)
    return {"count": len(builds), "builds": builds}


@router.get("/effective/{build_id}")
def get_effective_build(build_id: str, x_admin_token: str | None = Header(default=None)) -> dict[str, Any]:
    _require_admin(x_admin_token)
    return _load_build(build_id)


@router.get("/effective/{build_id}/artifact/{artifact}")
def get_effective_artifact(
    build_id: str,
    artifact: str,
    download: bool = False,
    x_admin_token: str | None = Header(default=None),
):
    _require_admin(x_admin_token)
    _load_build(build_id)
    mapping = {"markdown": "effective.md", "atomic": "effective.jsonl"}
    filename = mapping.get(artifact, "")
    if not filename:
        raise HTTPException(status_code=404, detail="Unknown effective artifact.")
    path = _build_dir(build_id) / filename
    if not path.exists():
        raise HTTPException(status_code=404, detail="Artifact not found.")
    if download:
        return FileResponse(path, filename=path.name)
    return PlainTextResponse(path.read_text(encoding="utf-8"), media_type="text/plain; charset=utf-8")


@router.post("/effective/{build_id}/upsert")
def publish_effective_policy(
    build_id: str,
    request: PublishEffectiveRequest,
    x_admin_token: str | None = Header(default=None),
) -> dict[str, Any]:
    _require_admin(x_admin_token)
    build = _load_build(build_id)
    path = _build_dir(build_id) / "effective.md"
    if not path.exists():
        raise HTTPException(status_code=409, detail="Effective Markdown is missing.")
    try:
        result = _upsert_effective_bytes(build, path.read_bytes(), request.replace_active_effective)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Effective policy upsert failed: {exc}") from exc
    build["published"] = True
    build["published_at"] = _utc_now()
    build["published_file"] = result.get("file")
    _save_build(build)
    return {"ok": True, "build": build, "upsert": result}
